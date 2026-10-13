#!/usr/bin/env python3
"""Build the editable report from the numbered Markdown source documents.

Requires python-docx and Pillow. PDF export and visual inspection are separate
steps; see the tooling README. Markdown sources remain the content authority.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import textwrap
from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
LINK = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')

def set_attr(element, name, value):
    element.set(qn(name), str(value))

def text_runs(paragraph, text):
    """Preserve inline code/bold and external links without raw Markdown tokens."""
    token = re.compile(r'(\[([^\]]+)\]\(([^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*)')
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(2):
            label, url = m.group(2), m.group(3)
            if url.startswith(('https://', 'http://')):
                rel = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                link = OxmlElement('w:hyperlink')
                set_attr(link, 'r:id', rel)
                run = OxmlElement('w:r')
                props = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); set_attr(color, 'w:val', '244B6C'); props.append(color)
                run.append(props)
                node = OxmlElement('w:t'); node.text = label; run.append(node)
                link.append(run); paragraph._p.append(link)
            else:
                paragraph.add_run(label)
        elif m.group(4):
            r = paragraph.add_run(m.group(4)); r.font.name = 'Courier New'; r.font.size = Pt(9)
        else:
            paragraph.add_run(m.group(5)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])

def split_row(line):
    return [x.strip().replace('\\|', '|') for x in re.split(r'(?<!\\)\|', line.strip().strip('|'))]

def format_table(doc, rows):
    columns = len(rows[0])
    if any(len(row) != columns for row in rows):
        raise ValueError('Inconsistent Markdown table column count')
    table = doc.add_table(rows=1, cols=columns)
    table.autofit = False
    usable = doc.sections[-1].page_width.inches - doc.sections[-1].left_margin.inches - doc.sections[-1].right_margin.inches
    scores = []
    for col in range(columns):
        values = [re.sub(r'[`*]', '', r[col]) for r in rows]
        scores.append(min(42, max(9, sum(len(x) for x in values) / len(values)) ** .75))
    if rows[0][0].lower() in ('id', 'measure', 'milestone', 'risk'):
        scores[0] = min(scores[0], 5 if rows[0][0].lower() == 'id' else 12)
    widths = [usable * score / sum(scores) for score in scores]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for index, values in enumerate(rows):
        cells = table.rows[0].cells if index == 0 else table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd'); set_attr(shading, 'w:fill', '23445C' if index == 0 else ('F1F4F6' if index % 2 == 0 else 'FFFFFF')); tcpr.append(shading)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + side); set_attr(margin, 'w:w', 80); set_attr(margin, 'w:type', 'dxa'); margins.append(margin)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side); set_attr(edge, 'w:val', 'single'); set_attr(edge, 'w:sz', 4); set_attr(edge, 'w:color', 'D9D9D9'); borders.append(edge)
            tcpr.append(borders)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            text_runs(p, value)
            for r in p.runs:
                r.font.size = Pt(9)
                if index == 0:
                    r.bold = True; r.font.color.rgb = RGBColor.from_string('FFFFFF')
            if columns > 2 and len(value) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        props = table.rows[index]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); props.append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); props.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def section_settings(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)
    section.footer_distance = Inches(.3)

def add_figure(doc, image_path, caption, description='', chapter_title=None):
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    landscape = pixel_width / pixel_height > 1.2
    section_settings(doc.add_section(WD_SECTION_START.NEW_PAGE), landscape)
    width_limit = 9.5 if landscape else 7
    height_limit = (5.6 if landscape else 7.9) - (1.1 if chapter_title else 0)
    scale = min(width_limit / pixel_width, height_limit / pixel_height)
    if chapter_title:
        doc.add_heading(chapter_title, 1)
    doc.add_heading(caption, 2)
    if description:
        text_runs(doc.add_paragraph(), description)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_picture(str(image_path), width=Inches(pixel_width * scale), height=Inches(pixel_height * scale))
    cp = doc.add_paragraph('Editable source: ' + image_path.name.replace('.png', '.mmd'), 'Caption')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(3)
    cp.paragraph_format.space_after = Pt(0)

def parse_diagrams(doc, md_path, figures):
    lines = md_path.read_text().splitlines()
    chapter = lines[0].lstrip('# ')
    intro = []
    title = None
    description = []
    first = True
    for line in lines[1:]:
        line = line.strip()
        if line.startswith('### '):
            title = line[4:]
            description = []
        elif line.startswith('## '):
            continue
        elif line.startswith('!['):
            m = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
            image_path = (md_path.parent / m.group(2)).resolve()
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            text = ' '.join((intro if first else []) + description)
            add_figure(doc, image_path, re.sub(r'\s*[—–]\s*', ' ', title or m.group(1)), text, chapter if first else None)
            figures.append(image_path.name)
            first = False
        elif line:
            (intro if title is None else description).append(line)

def parse_document(doc, md_path, figures):
    lines = md_path.read_text().splitlines()
    i = 0
    pending = []
    def flush():
        if pending:
            p = doc.add_paragraph(); text_runs(p, ' '.join(pending)); pending.clear()
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            flush(); i += 1; continue
        if line.startswith('```'):
            flush(); code = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            for code_line in code:
                wrapped = textwrap.wrap(code_line, width=88, subsequent_indent='    ', replace_whitespace=False, drop_whitespace=False, break_long_words=False, break_on_hyphens=False) or ['']
                for display_line in wrapped:
                    p = doc.add_paragraph(style='Code')
                    r = p.add_run(display_line)
                    r.font.name = 'Courier New'
                    r.font.size = Pt(9)
            if code:
                p.paragraph_format.space_after = Pt(6)
            i += 1; continue
        image = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if image:
            flush(); image_path = (md_path.parent / image.group(2)).resolve()
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            add_figure(doc, image_path, image.group(1))
            figures.append(image_path.name)
            i += 1; continue
        heading = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading:
            flush(); level = len(heading.group(1))
            label = re.sub(r'\s*[—–]\s*', ' ', heading.group(2))
            p = doc.add_heading(label, min(level, 3))
            if level == 1:
                p.paragraph_format.page_break_before = True
            i += 1; continue
        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|?[\s:|\-]+$', lines[i + 1].strip()):
            flush(); rows = [split_row(line)]; i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            format_table(doc, rows); continue
        bullet = re.match(r'^([-*]|\d+\.)\s+(.+)$', line)
        if bullet:
            flush(); style = 'List Bullet' if bullet.group(1) in ('-', '*') else 'List Number'
            text_runs(doc.add_paragraph(style=style), bullet.group(2)); i += 1; continue
        if line.startswith('> '):
            line = line[2:]
        if line == '---':
            flush(); i += 1; continue
        pending.append(line); i += 1
    flush()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--docs-dir', type=Path, default=ROOT / 'docs/preproduction')
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'output/preproduction')
    args = parser.parse_args()
    sources = sorted(args.docs_dir.glob('[0-9][0-9]-*.md'))
    if len(sources) != 8:
        raise ValueError(f'Expected eight source chapters, found {len(sources)}')
    args.output_dir.mkdir(parents=True, exist_ok=True)
    doc = Document(); section_settings(doc.sections[0])
    for name in ('Normal', 'Body Text', 'List Bullet', 'List Number'):
        style = doc.styles[name]; style.font.name = 'Calibri'; style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(5); style.paragraph_format.line_spacing = 1.07
    for name, size in [('Title', 28), ('Heading 1', 20), ('Heading 2', 14), ('Heading 3', 11)]:
        s = doc.styles[name]; s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = name != 'Title'; s.font.color.rgb = RGBColor(0,0,0)
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
    doc.styles['Subtitle'].font.name = 'Calibri'
    doc.styles['Subtitle'].font.size = Pt(18)
    doc.styles['Subtitle'].font.color.rgb = RGBColor(0,0,0)
    code = doc.styles.add_style('Code', 1)
    code.font.name = 'Courier New'; code.font.size = Pt(9)
    code.paragraph_format.space_after = Pt(0); code.paragraph_format.line_spacing = 1
    doc.styles['Caption'].font.size = Pt(9); doc.styles['Caption'].font.color.rgb = RGBColor(0,0,0)
    doc.core_properties.title = 'SpeedTower Preproduction Specification'
    doc.core_properties.subject = 'Game design, requirements, architecture, procedural worlds and validation'
    doc.core_properties.author = 'Group 7 planning package'
    doc.add_paragraph('SpeedTower', 'Title')
    doc.add_paragraph('Preproduction Specification', 'Subtitle')
    doc.add_paragraph('Group 7 | CS 4361.002')
    doc.add_paragraph('Prepared for Adrian Alvarez, Alexandru Istrate and Omar Sabry')
    doc.add_paragraph('Version 0.2 | October 13, 2026')
    doc.add_paragraph('SpeedTower is an endless, seeded voxel castle platformer with circumferential movement, independently rotating tower rings and stored or environmental boosts. This specification defines its initial gameplay, system contracts, graphics pipeline and implementation milestones.')
    doc.add_paragraph('The recommended design uses a custom cylindrical 2.5D controller with Three.js presentation and a generation worker. Movement tuning and performance targets require prototype validation. The package supplies editable source documents and diagrams for the team to maintain as development proceeds.')
    doc.add_heading('Document contents', 2)
    for source in sources:
        title = source.read_text().splitlines()[0].lstrip('# ').replace('—', '').strip()
        doc.add_paragraph(title, 'List Bullet')
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('SpeedTower  |  ')
    field = OxmlElement('w:fldSimple'); set_attr(field, 'w:instr', 'PAGE'); footer._p.append(field)
    figures = []
    previous_diagrams = False
    for source in sources:
        if source.name.startswith('06-'):
            parse_diagrams(doc, source, figures)
            previous_diagrams = True
        else:
            if previous_diagrams:
                section_settings(doc.add_section(WD_SECTION_START.NEW_PAGE), False)
            parse_document(doc, source, figures)
            previous_diagrams = False
    # Remove inherited Word title rules and paragraph borders from the report.
    for container in [doc._element, doc.styles.element]:
        for element in container.iter(qn('w:pBdr')):
            element.getparent().remove(element)
    target = args.output_dir / 'SpeedTower-Preproduction-Specification.docx'
    doc.save(target)
    (args.output_dir / 'report-manifest.json').write_text(json.dumps({'sources':[{'file':x.name, 'sha256':hashlib.sha256(x.read_bytes()).hexdigest()} for x in sources], 'figures':figures, 'figureHashes':{name:hashlib.sha256((args.docs_dir/'diagrams/rendered'/name).read_bytes()).hexdigest() for name in figures}, 'docx':target.name}, indent=2) + '\n')
    print(target)

if __name__ == '__main__':
    main()
